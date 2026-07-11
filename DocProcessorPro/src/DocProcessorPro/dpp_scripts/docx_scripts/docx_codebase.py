"""Renders a large-model case summary (see case_summary.py) into a .docx file."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject


def build_summary_docx(case_summary: dict, case_label: str, output_path: "str | Path") -> None:
    """Writes a case-summary .docx to output_path from the large model's
    structured output: {"records_reviewed": [...], "provider_chronology":
    [{"provider_name", "location_name", "date_range", "entries": [{"date",
    "record_type", "summary"}]}], "narrative_summary": "..."}.
    """
    doc = Document()
    doc.add_heading(case_label or "Medical Record Summary", level=0)

    records_reviewed = case_summary.get("records_reviewed") or []
    if records_reviewed:
        doc.add_heading("Records Reviewed", level=1)
        for line in records_reviewed:
            doc.add_paragraph(_strip_leading_number(line), style="List Number")

    provider_chronology = case_summary.get("provider_chronology") or []
    if provider_chronology:
        doc.add_heading("Provider Chronology", level=1)
        for provider in provider_chronology:
            _add_provider_section(doc, provider)

    narrative_summary = case_summary.get("narrative_summary")
    if narrative_summary:
        doc.add_heading("Narrative Summary", level=1)
        doc.add_paragraph(narrative_summary)

    out_path = Path(output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def _add_provider_section(doc: DocumentObject, provider: dict) -> None:
    provider_name = provider.get("provider_name") or "Unknown Provider"
    location_name = provider.get("location_name")
    date_range = provider.get("date_range")

    heading_text = provider_name
    if location_name and location_name != provider_name:
        heading_text += f" — {location_name}"
    doc.add_heading(heading_text, level=2)

    if date_range:
        date_para = doc.add_paragraph()
        date_run = date_para.add_run(f"Dates of service: {date_range}")
        date_run.italic = True

    entries = provider.get("entries") or []
    if not entries:
        return

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = "Date"
    header[1].text = "Record Type"
    header[2].text = "Summary"
    for entry in entries:
        row = table.add_row().cells
        row[0].text = str(entry.get("date") or "")
        row[1].text = str(entry.get("record_type") or "")
        row[2].text = str(entry.get("summary") or "")


def _strip_leading_number(line: str) -> str:
    """build_records_reviewed_list() already numbers each line ('1. ...');
    Word's 'List Number' paragraph style adds its own numbering, so strip the
    literal digit prefix to avoid a doubled '1. 1. ...' look."""
    stripped = line.lstrip()
    dot = stripped.find(". ")
    if dot != -1 and stripped[:dot].isdigit():
        return stripped[dot + 2 :]
    return stripped
